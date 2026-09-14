"""
Produces a revised taxonomy workbook (vNEXT) from an original vPREV workbook
plus a JSON revision spec.

WHY A SPEC FILE RATHER THAN HAND-EDITING THE XLSX
The taxonomy workbooks are the human-facing artefact, but a hand-edited xlsx
loses the audit trail: six months later nobody can tell which rows came from
the original research and which came from a real pipeline run, or why. This
script keeps the original workbook untouched, applies an explicit, reviewable
list of changes, and writes a Change Log sheet naming the evidence for each
one. Re-running it reproduces the same output.

SPEC SHAPE (JSON)
{
  "regulator": "MERC",
  "version": "v1.1",
  "previous_version": "v1.0",
  "run_note": "one paragraph: what run produced these changes",
  "add_tags":      [{"facet","name","definition","notes"}],
  "revise_tags":   [{"facet","name","definition"?,"notes"?,"status"?}],
  "deprecate_tags":[{"facet","name","reason"}],
  "add_rules":     [{"rule","explanation","example"}],
  "revise_rules":  [{"match","rule","explanation","example"}],
  "add_notes":     [{"topic","note"}],
  "dataset":       [{"date","title","instrument","subject","notes"}],
  "changelog":     [{"change","evidence"}]
}

Every field is optional except regulator/version. A tag named in revise_tags or
deprecate_tags that does not exist in the original workbook is a hard error --
that means the spec and the workbook have drifted, and silently ignoring it
would produce a workbook that claims a change it did not make.

Usage:
  python scripts/revise_taxonomy_workbook.py \
      --source "MERC_Regulatory_Taxonomy_v1_0.xlsx" \
      --spec   scripts/.audit/merc-revision.json \
      --out    MERC_Regulatory_Taxonomy_v1_1.xlsx

Both the vPREV source workbook and the vNEXT output live in the repo root, so
this is reproducible on any checkout -- do not point --source at a path outside
the repository.
"""

from __future__ import annotations

import argparse
import json
import sys
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ADDED_FILL = PatternFill("solid", fgColor="E2EFDA")     # green: new in this version
REVISED_FILL = PatternFill("solid", fgColor="FFF2CC")   # amber: changed
DEPRECATED_FILL = PatternFill("solid", fgColor="FCE4E4")  # red: deprecated
TITLE_FONT = Font(bold=True, size=13)
WRAP = Alignment(wrap_text=True, vertical="top")


def _rows(ws) -> list[list]:
    return [[c.value for c in row] for row in ws.iter_rows()]


def _find_header_row(rows: list[list], first_cell: str) -> int:
    for idx, row in enumerate(rows):
        if row and str(row[0]).strip() == first_cell:
            return idx
    raise ValueError(f"could not find a header row starting with {first_cell!r}")


def _style_sheet(ws, header_row: int, widths: list[int]) -> None:
    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = WRAP
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = WRAP
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def build_taxonomy_sheet(wb: Workbook, source_rows: list[list], spec: dict) -> None:
    ws = wb.create_sheet("Taxonomy")
    version = spec["version"]
    previous = spec.get("previous_version", "v1.0")

    header_idx = _find_header_row(source_rows, "Regulator")
    header = source_rows[header_idx]
    preamble = source_rows[:header_idx]

    ws.append([f"{spec['regulator']} Taxonomy — {version}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    for row in preamble[1:]:
        ws.append([row[0] if row else ""])
    ws.append([spec.get("run_note", "")])
    ws.append([])
    header_row_number = ws.max_row + 1
    ws.append(list(header))

    revise_by_key = {(r["facet"], r["name"]): r for r in spec.get("revise_tags", [])}
    deprecate_by_key = {(r["facet"], r["name"]): r for r in spec.get("deprecate_tags", [])}
    seen_keys: set[tuple[str, str]] = set()

    for row in source_rows[header_idx + 1:]:
        if not row or not row[1]:
            continue
        row = list(row) + [None] * (len(header) - len(row))
        facet, name = str(row[1]).strip(), str(row[2]).strip()
        key = (facet, name)
        seen_keys.add(key)
        fill = None

        if key in revise_by_key:
            change = revise_by_key[key]
            if "definition" in change:
                row[4] = change["definition"]
            if "notes" in change:
                row[7] = change["notes"]
            if "status" in change:
                row[5] = change["status"]
            row[6] = f"{row[6]} (revised {version})" if row[6] else version
            fill = REVISED_FILL
        elif key in deprecate_by_key:
            row[5] = "Deprecated"
            row[7] = f"DEPRECATED in {version}. {deprecate_by_key[key]['reason']}"
            fill = DEPRECATED_FILL

        ws.append(row)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    for change in spec.get("add_tags", []):
        ws.append(
            [
                spec["regulator"],
                change["facet"],
                change["name"],
                change.get("parent", ""),
                change["definition"],
                change.get("status", "Active"),
                version,
                change.get("notes", ""),
            ]
        )
        for cell in ws[ws.max_row]:
            cell.fill = ADDED_FILL

    missing = (set(revise_by_key) | set(deprecate_by_key)) - seen_keys
    if missing:
        raise SystemExit(
            f"spec names {len(missing)} tag(s) not present in the {previous} workbook: "
            + ", ".join(f"{f}/{n}" for f, n in sorted(missing))
        )

    _style_sheet(ws, header_row_number, [12, 16, 34, 14, 62, 12, 12, 70])


def build_guide_sheet(wb: Workbook, source_rows: list[list], spec: dict) -> None:
    ws = wb.create_sheet("Tagging Guide")
    header_idx = _find_header_row(source_rows, "Rule")
    ws.append([f"{spec['regulator']} — Tagging Guide {spec['version']}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([])
    header_row_number = ws.max_row + 1
    ws.append(list(source_rows[header_idx]))

    revise_rules = spec.get("revise_rules", [])
    matched: set[int] = set()
    for row in source_rows[header_idx + 1:]:
        if not row or not row[0]:
            continue
        row = list(row) + [None] * (3 - len(row))
        fill = None
        for index, change in enumerate(revise_rules):
            if change["match"].lower() in str(row[0]).lower():
                row[0] = change.get("rule", row[0])
                row[1] = change.get("explanation", row[1])
                row[2] = change.get("example", row[2])
                matched.add(index)
                fill = REVISED_FILL
                break
        ws.append(row)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    unmatched = [c["match"] for i, c in enumerate(revise_rules) if i not in matched]
    if unmatched:
        raise SystemExit(f"revise_rules matched nothing for: {unmatched}")

    for change in spec.get("add_rules", []):
        ws.append([change["rule"], change["explanation"], change.get("example", "")])
        for cell in ws[ws.max_row]:
            cell.fill = ADDED_FILL

    _style_sheet(ws, header_row_number, [62, 84, 62])


def build_dataset_sheet(wb: Workbook, spec: dict) -> None:
    ws = wb.create_sheet("Dataset")
    ws.append([f"{spec['regulator']} — Real Classified Sample {spec['version']}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([spec.get("dataset_note", "")])
    ws.append([])
    header_row_number = ws.max_row + 1
    ws.append(["Regulator", "Date", "Title", "Instrument Type", "Subject", "Status", "Notes"])
    for item in spec.get("dataset", []):
        ws.append(
            [
                spec["regulator"],
                item.get("date", ""),
                item.get("title", ""),
                item.get("instrument", ""),
                item.get("subject", ""),
                item.get("status", ""),
                item.get("notes", ""),
            ]
        )
    _style_sheet(ws, header_row_number, [12, 13, 74, 26, 34, 24, 62])


def build_schema_sheet(wb: Workbook, source_rows: list[list], spec: dict) -> None:
    ws = wb.create_sheet("Schema Note")
    header_idx = _find_header_row(source_rows, "Topic")
    ws.append([f"{spec['regulator']} — Schema Notes {spec['version']}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([])
    header_row_number = ws.max_row + 1
    ws.append(list(source_rows[header_idx]))
    for row in source_rows[header_idx + 1:]:
        if row and row[0]:
            ws.append(list(row))
    for note in spec.get("add_notes", []):
        ws.append([note["topic"], note["note"]])
        for cell in ws[ws.max_row]:
            cell.fill = ADDED_FILL
    _style_sheet(ws, header_row_number, [46, 120])


def build_changelog_sheet(wb: Workbook, spec: dict) -> None:
    ws = wb.create_sheet("Change Log")
    ws.append([f"{spec['regulator']} — {spec.get('previous_version', 'v1.0')} to {spec['version']}"])
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([spec.get("run_note", "")])
    ws.append([])
    header_row_number = ws.max_row + 1
    ws.append(["Change", "Evidence from the real run"])
    for item in spec.get("changelog", []):
        ws.append([item["change"], item["evidence"]])
    _style_sheet(ws, header_row_number, [70, 120])


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply a revision spec to a taxonomy workbook")
    parser.add_argument("--source", required=True)
    parser.add_argument("--spec", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    with open(args.spec, encoding="utf-8") as fh:
        spec = json.load(fh)

    source = load_workbook(args.source)
    taxonomy_rows = _rows(source["Taxonomy"])
    guide_rows = _rows(source["Tagging Guide"])
    schema_rows = _rows(source["Schema Note"])

    wb = Workbook()
    wb.remove(wb.active)
    build_taxonomy_sheet(wb, taxonomy_rows, spec)
    build_guide_sheet(wb, guide_rows, spec)
    build_dataset_sheet(wb, spec)
    build_schema_sheet(wb, schema_rows, spec)
    build_changelog_sheet(wb, spec)
    wb.save(args.out)

    added = len(spec.get("add_tags", []))
    revised = len(spec.get("revise_tags", []))
    deprecated = len(spec.get("deprecate_tags", []))
    print(f"Wrote {args.out}")
    print(f"  tags added {added}, revised {revised}, deprecated {deprecated}")
    print(f"  rules added {len(spec.get('add_rules', []))}, revised {len(spec.get('revise_rules', []))}")
    print(f"  schema notes added {len(spec.get('add_notes', []))}")
    print(f"  dataset rows (real, from the run) {len(spec.get('dataset', []))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
